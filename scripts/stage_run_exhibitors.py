"""Turn a research run into a staged exhibitor file conference_intake can land.

    python3 scripts/stage_run_exhibitors.py --placements /tmp/placements.json

WHY THIS EXISTS RATHER THAN A DIRECT WRITE. A research run hands over names,
websites and a read of who buys. It does not hand over a company record: sector
and category are a judgement, an ATS has to be discovered and verified, and
`validate()` refuses a hollow row - correctly. conference_intake already owns
the three-way split the owner ruled on (already on file -> extend the
exhibited-at tag; new and not govtech -> suppliers.json; new and govtech ->
govtech_candidates.json for research). So the run's job is to produce the input
that split expects, and nothing more.

THE PLACEMENT IS EVIDENCE, NOT A RULING. `is_govtech` here carries the reason
and the verifier's verdict beside it, so a person reading the candidates queue
sees what the machine thought AND what the adversarial check said about it.
A placement two readers split on is worth more than either one's confidence.
"""
from __future__ import annotations

import argparse
import json
import pathlib
import re
import sys

ROOT = pathlib.Path(__file__).resolve().parent.parent
DATA = ROOT / "data"

# The run, its workbook, and the conference tag its names were read off. A run
# with no event (the Gov IT list arrived with no show named) carries None and
# is staged without a tag - conference_intake refuses a tag the catalogue never
# issued, which is exactly right: a name invented here could not be read back.
RUNS = {
    "stormcon": {
        "xlsx": "/Users/wyethwest/Downloads/SLED-stormcon-74.xlsx",
        "name_col": "Name", "site_col": "Website", "desc_col": "Description",
        "event_tag": "StormCon 2026",
        "source_url": "https://scon26.mapyourshow.com/",
        "source_note": (
            "The StormCon 2026 expo exhibitor page ('74 Results'), pasted by "
            "the owner on 2026-09-10 and then researched one company at a "
            "time. scon26.mapyourshow.com refuses fetches, so no name could be "
            "confirmed against the show's own URL - the list's provenance is "
            "the owner's paste, and every website beside it was confirmed from "
            "the company's own page."),
    },
    "govit": {
        "xlsx": "/Users/wyethwest/Downloads/SLED-gov-it-cyber-35.xlsx",
        "name_col": "Name (as on site)", "site_col": "Website",
        "desc_col": "Description",
        "event_tag": None,
        "source_url": None,
        "source_note": (
            "35 names pasted by the owner on 2026-09-11 with NO event named. "
            "The run guessed a state/local gov-IT or GMIS-style exhibitor list "
            "and that guess is not evidence, so no conference tag is claimed "
            "here. Staged as a plain research list."),
    },
}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--placements", required=True,
                    help="the placement pass output: [{name, outcome, sector, "
                         "category, reason, confidence, verify:{...}}]")
    ap.add_argument("--out-dir", default=str(DATA))
    a = ap.parse_args()

    placed = json.loads(pathlib.Path(a.placements).read_text())
    if isinstance(placed, dict):
        placed = placed.get("placements", [])
    by_name = {p["name"]: p for p in placed}

    import openpyxl
    for key, spec in RUNS.items():
        ws = openpyxl.load_workbook(spec["xlsx"], read_only=True)["Companies"]
        hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        H = {h: i for i, h in enumerate(hdr) if h}
        exhibitors = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            name = r[H[spec["name_col"]]]
            if not name:
                continue
            p = by_name.get(str(name)) or {}
            v = p.get("verify") or {}
            outcome = p.get("outcome")
            # not_a_company NEVER becomes a row. An agency, an association or a
            # product line is not an employer, and filing one as a supplier
            # puts a name on the map that can never be researched into anything.
            if outcome == "not_a_company":
                continue
            exhibitors.append({
                "name": str(name),
                "website": r[H[spec["site_col"]]],
                "description": (str(r[H[spec["desc_col"]]])[:300]
                                if spec["desc_col"] in H
                                and r[H[spec["desc_col"]]] else None),
                # A NULL HERE IS HONEST. An unplaced name was never judged, and
                # conference_intake treats null as "not judged" rather than as
                # a no - which is the difference between a catalogued supplier
                # and a vendor nobody ever looked at again.
                "is_govtech": (True if outcome == "vendor"
                               else False if outcome == "supplier" else None),
                "is_govtech_why": p.get("reason") or "not judged in this pass",
                # `vertical` ONLY FOR A CANDIDATE. conference_intake writes a
                # supplier's description as `vertical or description`, so
                # putting the category here would give all 78 suppliers a
                # description reading "Suppliers & Services" - a field that
                # says nothing, on every row, forever. A candidate's vertical
                # IS the category, because the research pass reads it as the
                # proposed placement.
                "vertical": (p.get("category")
                             if p.get("outcome") == "vendor" else None),
                "sector": p.get("sector"),
                "confidence": p.get("confidence"),
                "verifier": v.get("verdict"),
                "verifier_why": v.get("why"),
            })
        doc = {
            "event_tag": spec["event_tag"],
            "conference": (spec["event_tag"] or f"{key} research list"),
            "found": True,
            # NOT A PAGE SCRAPE, so it cannot be graded the way a sweep is.
            # sweep_exhibitors grades a capture `menu` when its names read as
            # a site's own navigation; that failure is impossible here, because
            # every name was researched to its own website one at a time. The
            # grade says so rather than borrowing a word that would mean the
            # detector had run.
            "quality": "researched",
            "source_url": spec["source_url"],
            "source_note": spec["source_note"],
            "exhibitors": exhibitors,
        }
        slug = re.sub(r"[^A-Za-z0-9]+", "_",
                      spec["event_tag"] or f"run_{key}")
        out = pathlib.Path(a.out_dir) / f"exhibitors_{slug}.json"
        out.write_text(json.dumps(doc, indent=1, ensure_ascii=False) + "\n")
        judged = sum(1 for e in exhibitors if e["is_govtech"] is not None)
        vendors = sum(1 for e in exhibitors if e["is_govtech"] is True)
        print(f"{key}: {len(exhibitors)} exhibitors staged -> {out.name}")
        print(f"   {judged} judged, {vendors} read as govtech vendors, "
              f"{judged - vendors} as suppliers, "
              f"{len(exhibitors) - judged} unjudged")
        dis = [e for e in exhibitors
               if e.get("verifier") and e["verifier"] != "agree"]
        if dis:
            print(f"   {len(dis)} DISPUTED by the adversarial check - these are "
                  f"what a person should read first:")
            for e in dis[:10]:
                print(f"      {e['name'][:36]:36} {e['verifier']}: "
                      f"{str(e['verifier_why'])[:70]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
